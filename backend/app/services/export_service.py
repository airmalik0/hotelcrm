import csv
import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


class ExportService:
    def __init__(self) -> None:
        self.header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        self.header_font = Font(color="FFFFFF", bold=True)
        self.header_alignment = Alignment(horizontal="center", vertical="center")

    def export_to_csv(self, data: list[dict[str, Any]], filename: str = "export.csv") -> bytes:
        """Export data to CSV format."""
        if not data:
            return b""

        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=data[0].keys())
        writer.writeheader()
        writer.writerows(data)

        return output.getvalue().encode("utf-8")

    def export_to_excel(
        self,
        data: list[dict[str, Any]] | dict[str, list[dict[str, Any]]],
        filename: str = "export.xlsx",
        include_summary: bool = True,
    ) -> bytes:
        """Export data to Excel format with formatting."""
        wb = Workbook()

        # Handle single sheet or multiple sheets
        if isinstance(data, list):
            sheets_data = {"Report": data}
        else:
            sheets_data = data

        # Remove default sheet if we have data
        if sheets_data and wb.active:
            wb.remove(wb.active)

        for sheet_name, sheet_data in sheets_data.items():
            ws = wb.create_sheet(title=sheet_name[:31])  # Excel sheet name limit

            if not sheet_data:
                ws.append(["No data available"])
                continue

            # Write headers
            headers = list(sheet_data[0].keys())
            ws.append(headers)

            # Format headers
            for col_idx, _ in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = self.header_fill
                cell.font = self.header_font
                cell.alignment = self.header_alignment

            # Write data
            for row_data in sheet_data:
                row = []
                for header in headers:
                    value = row_data.get(header)
                    # Format datetime objects
                    if isinstance(value, datetime):
                        value = value.strftime("%Y-%m-%d %H:%M:%S")
                    row.append(value)
                ws.append(row)

            # Auto-adjust column widths
            for col_idx, column_cells in enumerate(ws.columns, 1):
                max_length = 0
                column_letter = get_column_letter(col_idx)

                for cell in column_cells:
                    try:
                        if cell.value:
                            max_length = max(max_length, len(str(cell.value)))
                    except Exception:
                        pass

                adjusted_width = min(max_length + 2, 50)  # Cap at 50 characters
                ws.column_dimensions[column_letter].width = adjusted_width

            # Add summary if requested
            if include_summary and sheet_data:
                self._add_summary_row(ws, sheet_data, headers)

        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.read()

    def _add_summary_row(self, ws: Any, data: list[dict[str, Any]], headers: list[str]) -> None:
        """Add summary row with formulas for numeric columns."""
        summary_row = []

        for header in headers:
            # Check if column contains numeric data
            numeric_values = []
            for row in data:
                value = row.get(header)
                if isinstance(value, int | float):
                    numeric_values.append(value)

            if numeric_values and len(numeric_values) == len(data):
                # Add SUM formula for fully numeric columns
                col_letter = get_column_letter(headers.index(header) + 1)
                start_row = 2  # Data starts from row 2 (after headers)
                end_row = len(data) + 1
                summary_row.append(f"=SUM({col_letter}{start_row}:{col_letter}{end_row})")
            else:
                summary_row.append("" if header != headers[0] else "TOTAL")

        if any(cell for cell in summary_row[1:]):  # If there are any formulas
            ws.append(summary_row)
            # Format summary row
            last_row = ws.max_row
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=last_row, column=col_idx)
                cell.font = Font(bold=True)

    def export_occupancy_report(
        self, occupancy_data: list[dict[str, Any]], format: str = "excel"
    ) -> bytes:
        """Export occupancy report with specific formatting."""
        if format == "csv":
            return self.export_to_csv(occupancy_data)

        # Excel with multiple analysis sheets
        sheets_data = {
            "Occupancy Overview": occupancy_data,
        }

        # Add calculated metrics sheet
        if occupancy_data:
            metrics = self._calculate_occupancy_metrics(occupancy_data)
            sheets_data["Metrics"] = metrics

        return self.export_to_excel(sheets_data)

    def _calculate_occupancy_metrics(self, data: list[dict[str, Any]]) -> list[dict[str, Any]]:
        """Calculate occupancy metrics for the report."""
        if not data:
            return []

        total_bookings = sum(row.get("total_bookings", 0) for row in data)
        avg_occupancy = sum(row.get("occupancy_rate", 0) for row in data) / len(data)
        max_occupancy = max(row.get("occupancy_rate", 0) for row in data)
        min_occupancy = min(row.get("occupancy_rate", 0) for row in data)

        return [
            {"Metric": "Total Bookings", "Value": total_bookings},
            {"Metric": "Average Occupancy Rate", "Value": f"{avg_occupancy:.2f}%"},
            {"Metric": "Maximum Occupancy Rate", "Value": f"{max_occupancy:.2f}%"},
            {"Metric": "Minimum Occupancy Rate", "Value": f"{min_occupancy:.2f}%"},
            {"Metric": "Reporting Periods", "Value": len(data)},
        ]

    def export_revenue_report(
        self, revenue_data: list[dict[str, Any]], format: str = "excel"
    ) -> bytes:
        """Export revenue report with financial formatting."""
        if format == "csv":
            return self.export_to_csv(revenue_data)

        # Format currency values
        formatted_data = []
        for row in revenue_data:
            formatted_row = row.copy()
            for key in ["total_revenue", "avg_booking_value"]:
                if key in formatted_row and formatted_row[key] is not None:
                    formatted_row[key] = f"${formatted_row[key]:,.2f}"
            formatted_data.append(formatted_row)

        return self.export_to_excel({"Revenue Report": formatted_data})

    def export_customer_report(
        self, customer_data: list[dict[str, Any]], format: str = "excel"
    ) -> bytes:
        """Export customer report with contact information."""
        if format == "csv":
            return self.export_to_csv(customer_data)

        # Format for Excel with customer insights
        formatted_data = []
        for row in customer_data:
            formatted_row = row.copy()
            if "total_revenue" in formatted_row and formatted_row["total_revenue"] is not None:
                formatted_row["total_revenue"] = f"${formatted_row['total_revenue']:,.2f}"
            formatted_data.append(formatted_row)

        return self.export_to_excel({"Top Customers": formatted_data})

    def export_payment_methods_report(
        self, payment_data: list[dict[str, Any]], format: str = "excel"
    ) -> bytes:
        """Export payment methods distribution report."""
        if format == "csv":
            # Flatten nested payment method data for CSV
            flattened_data = []
            for period in payment_data:
                row = {
                    "Period": period["period_label"],
                    "Cash Count": period["cash"]["count"],
                    "Cash Amount": period["cash"]["amount"],
                    "Cash %": period["cash"].get("percentage", 0),
                    "Terminal Count": period["terminal"]["count"],
                    "Terminal Amount": period["terminal"]["amount"],
                    "Terminal %": period["terminal"].get("percentage", 0),
                    "Transfer Count": period["transfer"]["count"],
                    "Transfer Amount": period["transfer"]["amount"],
                    "Transfer %": period["transfer"].get("percentage", 0),
                    "Total Bookings": period["total_bookings"],
                    "Total Amount": period["total_amount"],
                }
                flattened_data.append(row)
            return self.export_to_csv(flattened_data)

        # Excel with formatted payment data
        return self.export_to_excel({"Payment Methods": payment_data})

    def export_geographic_report(
        self, geographic_data: list[dict[str, Any]], format: str = "excel"
    ) -> bytes:
        """Export geographic analysis report."""
        if format == "csv":
            return self.export_to_csv(geographic_data)

        # Format currency values
        formatted_data = []
        for row in geographic_data:
            formatted_row = row.copy()
            for key in ["total_revenue", "avg_booking_value"]:
                if key in formatted_row and formatted_row[key] is not None:
                    formatted_row[key] = f"${formatted_row[key]:,.2f}"
            formatted_data.append(formatted_row)

        # Create summary statistics
        if geographic_data:
            total_customers = sum(row.get("customer_count", 0) for row in geographic_data)
            total_bookings = sum(row.get("booking_count", 0) for row in geographic_data)
            total_revenue = sum(row.get("total_revenue", 0) for row in geographic_data)

            summary = [
                {"Metric": "Total Districts", "Value": len(geographic_data)},
                {"Metric": "Total Customers", "Value": total_customers},
                {"Metric": "Total Bookings", "Value": total_bookings},
                {"Metric": "Total Revenue", "Value": f"${total_revenue:,.2f}"},
            ]

            return self.export_to_excel({
                "Geographic Analysis": formatted_data,
                "Summary": summary,
            })

        return self.export_to_excel({"Geographic Analysis": formatted_data})
